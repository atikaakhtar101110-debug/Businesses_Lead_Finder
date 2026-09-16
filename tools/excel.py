"""
Exports a list of Lead objects to a formatted .xlsx workbook.
"""
import os
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment
from openpyxl.utils import get_column_letter
from utils.logger import get_logger

logger = get_logger("tools.excel")

HEADER_FILL = PatternFill(start_color="1F4E78", end_color="1F4E78", fill_type="solid")
HEADER_FONT = Font(color="FFFFFF", bold=True)


class ExcelExporter:
    def export_leads(self, leads: list, path: str) -> str:
        os.makedirs(os.path.dirname(path) or ".", exist_ok=True)

        wb = Workbook()
        ws = wb.active
        ws.title = "Leads"

        if not leads:
            ws.append(["No leads found"])
            wb.save(path)
            return path

        rows = [lead.to_row() if hasattr(lead, "to_row") else lead for lead in leads]
        headers = list(rows[0].keys())
        ws.append(headers)

        for col_idx, _ in enumerate(headers, start=1):
            cell = ws.cell(row=1, column=col_idx)
            cell.fill = HEADER_FILL
            cell.font = HEADER_FONT
            cell.alignment = Alignment(horizontal="center", vertical="center")

        # Sort rows by Score descending if present
        if "Score" in headers:
            rows.sort(key=lambda r: r.get("Score", 0) or 0, reverse=True)

        for row in rows:
            ws.append([row.get(h, "") for h in headers])

        # Auto-size columns (capped so long descriptions don't blow out the sheet)
        for col_idx, header in enumerate(headers, start=1):
            max_len = max(
                [len(str(header))] + [len(str(r.get(header, ""))) for r in rows]
            )
            ws.column_dimensions[get_column_letter(col_idx)].width = min(max(max_len + 2, 10), 45)

        ws.freeze_panes = "A2"
        wb.save(path)
        logger.info("Exported %d leads to %s", len(rows), path)
        return path
